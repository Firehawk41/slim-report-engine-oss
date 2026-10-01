"""End-to-end test of the actual CLI entry point (slim_report_engine.cli)
against a real (temp-file) SQLite database and the real blank Chemical
intake form -- the same contract a VBA trigger would rely on: exit code
0 + output file on success, exit code 1 + a clear stderr message on a
recognized data problem (e.g. an unresolved customer).
"""

import os
from datetime import date
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from infrastructure.database import Base
from slim_domain.domain.analysis.analysis_repository import _AnalysisRow, _FormAnalysisRow
from slim_domain.domain.chemical.chemical_repository import _ChemicalRow, _FormChemicalRow
from slim_domain.domain.customer.customer_repository import _CustomerRow, _FormCustomerRow
from slim_domain.domain.element.element_repository import _ElementRow

from slim_report_engine.cli import build_services, main, run
from slim_report_engine.lab_identity import LAB_HEADER_LINES
from slim_report_engine.reporting.presets import analyte_presets

_FORM_PATH = (
    Path(os.environ.get("PVB_BRIDGE_PATH", r"C:\Users\labuser\Python-VBA-Bridge"))
    / "ReportCreator Project"
    / "Doc-F31 EXAMPLE LAB Analytical Testing Request Form - Chemicals.xlsx"
)


def _seed(db_url: str) -> None:
    engine = create_engine(db_url)
    Base.metadata.create_all(engine)
    with Session(engine) as s:
        for symbol in analyte_presets.trace_elements_36():
            s.add(_ElementRow(element_symbol=symbol, element_name=symbol))

        customer = _CustomerRow(
            customer_name="Customer 130", street_address="123 Main St",
            city="Springfield", state="IL", postal_code="62701", country="USA",
        )
        s.add(customer)
        s.flush()
        s.add(_FormCustomerRow(
            form_customer="Customer 130", form_address="123 Main St", form_address_2="Springfield, IL",
            customer_id=customer.id,
        ))

        chemical = _ChemicalRow(chemical_name="Test Acid", metals_prep="N/A", silicon_prep="N/A", ions_prep="N/A")
        s.add(chemical)
        s.flush()
        s.add(_FormChemicalRow(form_name="Test Acid Matrix", chemical_id=chemical.id))

        analysis = _AnalysisRow(analysis_name="36 Elements", analysis_description="36 elements (by ICP-MS)")
        s.add(analysis)
        s.flush()
        s.add(_FormAnalysisRow(form_name="36 Elements", analysis_id=analysis.id))
        s.commit()


def _make_input_workbook(path: Path, customer_name: str = "Customer 130") -> None:
    wb = openpyxl.load_workbook(_FORM_PATH)
    ws = wb.active
    ws["C10"] = customer_name
    ws["C11"] = "123 Main St"
    ws["C12"] = "Springfield, IL"
    ws["P5"] = date(2026, 3, 5)
    ws["P7"] = "JT"
    ws["B22"] = "S-001"
    ws["C22"] = "Test Acid Matrix"
    ws["D22"] = "36 Elements"
    ws["K22"] = "Next Day"
    wb.save(path)


@pytest.mark.skipif(not _FORM_PATH.exists(), reason="real Chemical intake form not available in this environment")
def test_cli_success_writes_output_and_exits_zero(tmp_path, capsys):
    db_path = tmp_path / "test.db"
    db_url = f"sqlite:///{db_path}"
    _seed(db_url)

    input_path = tmp_path / "input.xlsx"
    _make_input_workbook(input_path)
    output_path = tmp_path / "output.xlsx"

    exit_code = main(["--db", db_url, str(input_path), "-o", str(output_path)])

    assert exit_code == 0
    assert output_path.exists()

    out, err = capsys.readouterr()
    assert str(output_path) in out
    assert err == ""

    wb = openpyxl.load_workbook(output_path)
    # Tab name is the chemical matrix (real system's own naming, see
    # submission_report_builder.py's SetChemicalSheetName port) -- not
    # the Sample ID.
    assert "Test Acid Matrix" in wb.sheetnames
    ws = wb["Test Acid Matrix"]
    assert ws.oddHeader.center.text == "Test Acid Matrix"
    assert ws.oddHeader.left.text == "\n".join(LAB_HEADER_LINES)
    assert "Customer 130" in ws.oddHeader.right.text
    assert ws["A1"].value == "Test Acid Matrix"
    # Confirmed real: a 5-row header preamble (title/blank/Notes:/2 note
    # lines) precedes the first section's own sample-ID echo row.
    assert ws["A3"].value == "Notes:"
    assert ws["A4"].value == "All data in units of ppb unless otherwise noted"
    assert ws["A5"].value == "Blue font indicates data at or below detection limits"
    # The standard "mmddyy-Chemical-Customer-SampleID" sample-ID string,
    # stamped into the section's own "Sample Identification:" row (row 6,
    # right after the preamble).
    assert ws["D6"].value == "030526-Test Acid Matrix-Customer 130-S-001"


@pytest.mark.skipif(not _FORM_PATH.exists(), reason="real Chemical intake form not available in this environment")
def test_cli_success_only_the_first_report_sheet_is_selected(tmp_path, capsys):
    """2026-09-29, direct user report: opening a real generated output
    showed both the first report tab AND the original Testing Request tab
    highlighted as "selected" (a confusing Excel multi-sheet-selection
    state) -- root cause confirmed via direct openpyxl inspection of the
    real Chemical intake template: its one sheet has sheet_view.
    tabSelected=True baked in, and neither wb.active nor wb.move_sheet
    ever clears that flag on a sheet once it's no longer active."""
    db_path = tmp_path / "test.db"
    db_url = f"sqlite:///{db_path}"
    _seed(db_url)

    input_path = tmp_path / "input.xlsx"
    _make_input_workbook(input_path)
    output_path = tmp_path / "output.xlsx"

    exit_code = main(["--db", db_url, str(input_path), "-o", str(output_path)])
    assert exit_code == 0

    wb = openpyxl.load_workbook(output_path)
    assert wb.active.title == "Test Acid Matrix"
    selected = [ws.title for ws in wb.worksheets if ws.sheet_view.tabSelected]
    assert selected == ["Test Acid Matrix"]


@pytest.mark.skipif(not _FORM_PATH.exists(), reason="real Chemical intake form not available in this environment")
def test_run_populates_the_optional_info_dict_with_the_report_gates(tmp_path):
    """2026-09-29, direct user request: Slim-LIMS's /generate-report
    response should surface "the big gates the code successfully clears"
    (customer/type/sample-count at parse, sheet-count at build, timing at
    each) for the VBA hook to log locally -- without changing run()'s own
    return type/contract, per its existing docstring. `info` defaults to
    None and is a no-op then (every other test in this file proves that
    implicitly, by never passing it); this test proves the populated
    path."""
    db_path = tmp_path / "test.db"
    db_url = f"sqlite:///{db_path}"
    _seed(db_url)
    services = build_services(db_url)

    input_path = tmp_path / "input.xlsx"
    _make_input_workbook(input_path)
    output_path = tmp_path / "output.xlsx"

    info: dict = {}
    result_path = run(input_path, output_path, services=services, info=info)

    assert result_path == output_path
    assert info["customer_name"] == "Customer 130"
    assert info["request_type"] == "CHEMICAL"
    assert info["sample_count"] == 1
    assert info["sheet_count"] == 1
    assert info["engine_version"] and info["engine_version"] != "unknown"
    assert info["domain_version"] and info["domain_version"] != "unknown"
    for key in ("parse_seconds", "build_seconds", "write_seconds"):
        assert isinstance(info[key], float)


@pytest.mark.skipif(not _FORM_PATH.exists(), reason="real Chemical intake form not available in this environment")
def test_cli_success_writes_a_report_generation_log_row(tmp_path, capsys):
    """End-to-end confirmation of the 2026-09-25 traceability feature: a
    successful run() writes one report_generation_log row, with the real
    technician initials pulled from the TR form's own "Initials:" cell."""
    db_path = tmp_path / "test.db"
    db_url = f"sqlite:///{db_path}"
    _seed(db_url)

    input_path = tmp_path / "input.xlsx"
    _make_input_workbook(input_path)
    output_path = tmp_path / "output.xlsx"

    exit_code = main(["--db", db_url, str(input_path), "-o", str(output_path)])
    assert exit_code == 0

    from slim_domain.domain.report_log.report_log_repository import _ReportGenerationLogRow

    with Session(create_engine(db_url)) as s:
        rows = s.query(_ReportGenerationLogRow).all()
    assert len(rows) == 1
    assert rows[0].customer_name == "Customer 130"
    assert rows[0].technician_initials == "JT"
    assert rows[0].input_path == str(input_path)
    assert rows[0].output_path == str(output_path)


@pytest.mark.skipif(not _FORM_PATH.exists(), reason="real Chemical intake form not available in this environment")
def test_cli_unresolved_customer_exits_one_with_clear_message(tmp_path, capsys):
    db_path = tmp_path / "test.db"
    db_url = f"sqlite:///{db_path}"
    _seed(db_url)

    input_path = tmp_path / "input.xlsx"
    _make_input_workbook(input_path, customer_name="Totally Unknown Customer")
    output_path = tmp_path / "output.xlsx"

    exit_code = main(["--db", db_url, str(input_path), "-o", str(output_path)])

    assert exit_code == 1
    assert not output_path.exists()

    out, err = capsys.readouterr()
    assert "Totally Unknown Customer" in err
    assert "database" in err.lower()


@pytest.mark.skipif(not _FORM_PATH.exists(), reason="real Chemical intake form not available in this environment")
def test_cli_blank_customer_field_exits_one_with_a_blank_specific_message(tmp_path, capsys):
    """2026-09-29 failure-path audit fix: a TR with NOTHING typed in the
    customer name field used to be routed through the same wording as a
    real, typo'd/unmapped customer name ("Unresolved customer: '' ... --
    this needs to be added to the database"), which reads as "add a blank
    customer to the database" -- nonsensical for a field the technician
    or customer just left empty. This is the exact "communicate well"
    scenario Jamie flagged from the legacy VBA macro (a blank required
    field silently producing no report/no notice) -- confirming here it
    both (a) still fails loudly (never silent) and (b) now says the right
    thing for a genuinely blank field specifically."""
    db_path = tmp_path / "test.db"
    db_url = f"sqlite:///{db_path}"
    _seed(db_url)

    input_path = tmp_path / "input.xlsx"
    _make_input_workbook(input_path, customer_name="")
    output_path = tmp_path / "output.xlsx"

    exit_code = main(["--db", db_url, str(input_path), "-o", str(output_path)])

    assert exit_code == 1
    assert not output_path.exists()
    out, err = capsys.readouterr()
    assert "No customer information was entered" in err
    assert "Unresolved customer" not in err


@pytest.mark.skipif(not _FORM_PATH.exists(), reason="real Chemical intake form not available in this environment")
def test_cli_blank_chemical_matrix_exits_one_with_a_blank_specific_message(tmp_path, capsys):
    """Same fix as the blank-customer test above, for the Chemical Matrix
    field -- this is the precise scenario Jamie described from the legacy
    macro: a TR with no chemical matrix entered."""
    db_path = tmp_path / "test.db"
    db_url = f"sqlite:///{db_path}"
    _seed(db_url)

    input_path = tmp_path / "input.xlsx"
    wb = openpyxl.load_workbook(_FORM_PATH)
    ws = wb.active
    ws["C10"] = "Customer 130"
    ws["C11"] = "123 Main St"
    ws["C12"] = "Springfield, IL"
    ws["P5"] = date(2026, 3, 5)
    ws["P7"] = "JT"
    ws["B22"] = "S-001"
    ws["C22"] = ""  # Chemical Matrix left blank
    ws["D22"] = "36 Elements"
    ws["K22"] = "Next Day"
    wb.save(input_path)
    output_path = tmp_path / "output.xlsx"

    exit_code = main(["--db", db_url, str(input_path), "-o", str(output_path)])

    assert exit_code == 1
    assert not output_path.exists()
    out, err = capsys.readouterr()
    assert "No chemical matrix was entered" in err
    assert "Unresolved chemical" not in err


@pytest.mark.skipif(not _FORM_PATH.exists(), reason="real Chemical intake form not available in this environment")
def test_cli_corrupted_input_file_exits_one_with_a_clear_message(tmp_path, capsys):
    """2026-09-29 failure-path audit fix: openpyxl.load_workbook() used to
    be completely unwrapped -- a corrupted download, wrong file type, or
    password-protected workbook raised a raw zipfile.BadZipFile straight
    through to main()'s generic exit-code-2/traceback path instead of a
    message a technician could act on."""
    input_path = tmp_path / "corrupted.xlsx"
    input_path.write_bytes(b"this is not a real xlsx file")

    exit_code = main([str(input_path)])

    assert exit_code == 1
    out, err = capsys.readouterr()
    assert "doesn't look like a valid, readable .xlsx workbook" in err
    assert "Traceback" not in err


def test_cli_missing_input_file_exits_one(tmp_path, capsys):
    exit_code = main([str(tmp_path / "does_not_exist.xlsx")])
    assert exit_code == 1
    out, err = capsys.readouterr()
    assert "not found" in err


@pytest.mark.skipif(not _FORM_PATH.exists(), reason="real Chemical intake form not available in this environment")
def test_cli_result_file_written_on_success(tmp_path):
    db_path = tmp_path / "test.db"
    db_url = f"sqlite:///{db_path}"
    _seed(db_url)

    input_path = tmp_path / "input.xlsx"
    _make_input_workbook(input_path)
    output_path = tmp_path / "output.xlsx"
    result_file = tmp_path / "result.txt"

    if not _FORM_PATH.exists():
        pytest.skip("real Chemical intake form not available in this environment")

    exit_code = main(["--db", db_url, str(input_path), "-o", str(output_path), "--result-file", str(result_file)])

    assert exit_code == 0
    code_line, message = result_file.read_text(encoding="utf-8").split("\n", 1)
    assert code_line == "0"
    assert message == str(output_path)


def test_cli_result_file_written_on_failure_no_tmp_file_left_behind(tmp_path):
    result_file = tmp_path / "result.txt"

    exit_code = main([str(tmp_path / "does_not_exist.xlsx"), "--result-file", str(result_file)])

    assert exit_code == 1
    code_line, message = result_file.read_text(encoding="utf-8").split("\n", 1)
    assert code_line == "1"
    assert "not found" in message
    assert not result_file.with_name(result_file.name + ".tmp").exists()
