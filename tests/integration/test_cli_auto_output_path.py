"""Covers run()'s output_path=None branch (the API server's path, once a
customer has a mapped Reports folder) -- specifically the destination
computed via ReportFolderService + report_filename(), which
tests/integration/test_cli.py never exercises (it only drives main(),
always with an explicit -o). See ReportAlreadyExistsError in cli.py for
why an existing file at that computed destination must stop the run
rather than being silently overwritten.
"""

import os
from datetime import date
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from infrastructure.database import Base
from slim_domain.domain.analysis.analysis_repository import _AnalysisRow, _FormAnalysisRow
from slim_domain.domain.analysis.analysis_service import AnalysisService
from slim_domain.domain.chemical.chemical_repository import _ChemicalRow, _FormChemicalRow
from slim_domain.domain.chemical.chemical_service import ChemicalService
from slim_domain.domain.customer.customer_repository import _CustomerRow, _FormCustomerRow
from slim_domain.domain.customer.customer_service import CustomerService
from slim_domain.domain.element.element_repository import _ElementRow
from slim_domain.domain.element.element_service import ElementService
from slim_domain.domain.report_folder.report_folder_service import ReportFolderService
from slim_domain.domain.report_log.report_log_service import ReportGenerationLogService
from slim_domain.domain.specification.specification_service import SpecificationService

from slim_report_engine.cli import ReportAlreadyExistsError, ReportEngineServices, run
from slim_report_engine.report_filename import report_filename
from slim_report_engine.reporting.presets import analyte_presets

_FORM_PATH = (
    Path(os.environ.get("PVB_BRIDGE_PATH", r"C:\Users\labuser\Python-VBA-Bridge"))
    / "ReportCreator Project"
    / "Doc-F31 EXAMPLE LAB Analytical Testing Request Form - Chemicals.xlsx"
)


def _seed_and_build_services(db_url: str, report_folder: Path) -> tuple[ReportEngineServices, int]:
    engine = create_engine(db_url)
    Base.metadata.create_all(engine)
    session = Session(engine)

    for symbol in analyte_presets.trace_elements_36():
        session.add(_ElementRow(element_symbol=symbol, element_name=symbol))

    customer = _CustomerRow(
        customer_name="Customer 130", street_address="123 Main St",
        city="Springfield", state="IL", postal_code="62701", country="USA",
    )
    session.add(customer)
    session.flush()
    session.add(_FormCustomerRow(
        form_customer="Customer 130", form_address="123 Main St", form_address_2="Springfield, IL",
        customer_id=customer.id,
    ))

    chemical = _ChemicalRow(chemical_name="Test Acid", metals_prep="N/A", silicon_prep="N/A", ions_prep="N/A")
    session.add(chemical)
    session.flush()
    session.add(_FormChemicalRow(form_name="Test Acid Matrix", chemical_id=chemical.id))

    analysis = _AnalysisRow(analysis_name="36 Elements", analysis_description="36 elements (by ICP-MS)")
    session.add(analysis)
    session.flush()
    session.add(_FormAnalysisRow(form_name="36 Elements", analysis_id=analysis.id))
    session.commit()
    customer_id = customer.id

    report_folder.mkdir(parents=True, exist_ok=True)
    services = ReportEngineServices(
        chemical_svc=ChemicalService(session),
        analysis_svc=AnalysisService(session),
        element_svc=ElementService(session),
        customer_svc=CustomerService(session),
        specification_svc=SpecificationService(session),
        report_folder_svc=ReportFolderService(session),
        report_log_svc=ReportGenerationLogService(session),
    )
    services.report_folder_svc.set_folder(customer_id, str(report_folder), created_by="JT")
    return services, customer_id


def _make_input_workbook(path: Path) -> None:
    wb = openpyxl.load_workbook(_FORM_PATH)
    ws = wb.active
    ws["C10"] = "Customer 130"
    ws["C11"] = "123 Main St"
    ws["C12"] = "Springfield, IL"
    ws["P5"] = date(2026, 3, 5)
    ws["B22"] = "S-001"
    ws["C22"] = "Test Acid Matrix"
    ws["D22"] = "36 Elements"
    ws["K22"] = "Next Day"
    wb.save(path)


pytestmark = pytest.mark.skipif(
    not _FORM_PATH.exists(), reason="real Chemical intake form not available in this environment"
)


def test_auto_output_path_writes_to_mapped_folder(tmp_path):
    db_url = f"sqlite:///{tmp_path / 'test.db'}"
    report_folder = tmp_path / "Reports" / "Customer 130"
    services, _ = _seed_and_build_services(db_url, report_folder)

    input_path = tmp_path / "090126 Customer 130 Widgets STD.xlsx"
    _make_input_workbook(input_path)

    result_path = run(input_path, None, db_url=None, services=services)

    assert result_path == report_folder / report_filename(input_path.name)
    assert result_path.exists()


def test_run_logs_version_customer_and_output_on_success(tmp_path, caplog):
    import logging

    db_url = f"sqlite:///{tmp_path / 'test.db'}"
    report_folder = tmp_path / "Reports" / "Customer 130"
    services, _ = _seed_and_build_services(db_url, report_folder)

    input_path = tmp_path / "090126 Customer 130 Widgets STD.xlsx"
    _make_input_workbook(input_path)

    with caplog.at_level(logging.INFO):
        result_path = run(input_path, None, db_url=None, services=services)

    success_lines = [r.getMessage() for r in caplog.records if "SUCCESS" in r.getMessage()]
    assert len(success_lines) == 1
    assert "Customer 130" in success_lines[0]
    assert input_path.name in success_lines[0]
    assert str(result_path) in success_lines[0]
    assert "slim-report-engine v" in success_lines[0]
    assert "slim-domain v" in success_lines[0]


def test_auto_output_path_refuses_to_overwrite_existing_file(tmp_path):
    db_url = f"sqlite:///{tmp_path / 'test.db'}"
    report_folder = tmp_path / "Reports" / "Customer 130"
    services, _ = _seed_and_build_services(db_url, report_folder)

    input_path = tmp_path / "090126 Customer 130 Widgets STD.xlsx"
    _make_input_workbook(input_path)

    first_result = run(input_path, None, db_url=None, services=services)
    original_bytes = first_result.read_bytes()

    with pytest.raises(ReportAlreadyExistsError) as excinfo:
        run(input_path, None, db_url=None, services=services)

    assert str(first_result) in str(excinfo.value)
    # The existing file (e.g. a technician's hand-edited final report)
    # must be left completely untouched, not partially overwritten.
    assert first_result.read_bytes() == original_bytes
