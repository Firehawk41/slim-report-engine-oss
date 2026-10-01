"""Covers run()'s output_path=None branch for SX5 specifically: North and
South are the same real "SX5" customer record (shared for invoicing), so
the destination folder can't come from customer_id alone the way every
other customer's does -- it has to also key on the TR filename's
North/South, mirroring sx5_location_from_filename (already used for
report content). See test_cli_auto_output_path.py for the equivalent
non-SX5, single-folder-per-customer coverage this deliberately doesn't
repeat.
"""

import os
from datetime import date
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from infrastructure.database import Base
from slim_domain.domain.analysis.analysis_repository import _AnalysisRow, _FormAnalysisRow
from slim_domain.domain.analysis.analysis_service import AnalysisService
from slim_domain.domain.chemical.chemical_repository import _ChemicalRow, _FormChemicalRow
from slim_domain.domain.chemical.chemical_service import ChemicalService
from slim_domain.domain.customer.customer_repository import _CustomerRow, _FormCustomerRow
from slim_domain.domain.customer.customer_service import CustomerService
from slim_domain.domain.element.element_repository import _ElementRow
from slim_domain.domain.element.element_service import ElementService
from slim_domain.domain.report_folder.report_folder_service import ReportFolderService
from slim_domain.domain.report_log.report_log_service import ReportGenerationLogService
from slim_domain.domain.specification.specification_repository import _SpecificationRow
from slim_domain.domain.specification.specification_service import SpecificationService

from slim_report_engine.cli import ReportEngineError, ReportEngineServices, run
from slim_report_engine.report_filename import report_filename
from slim_report_engine.reporting.presets import analyte_presets

_FORM_PATH = (
    Path(os.environ.get("PVB_BRIDGE_PATH", r"C:\Users\labuser\Python-VBA-Bridge"))
    / "ReportCreator Project"
    / "Doc-F31 EXAMPLE LAB Analytical Testing Request Form - Chemicals.xlsx"
)

pytestmark = pytest.mark.skipif(
    not _FORM_PATH.exists(), reason="real Chemical intake form not available in this environment"
)


def _seed_sx5_services(db_url: str, north_folder: Path, south_folder: Path) -> ReportEngineServices:
    engine = create_engine(db_url)
    Base.metadata.create_all(engine)
    session = Session(engine)

    for symbol in analyte_presets.trace_elements_36():
        session.add(_ElementRow(element_symbol=symbol, element_name=symbol))

    customer = _CustomerRow(
        customer_name="SX5", street_address="1 Fab Way",
        city="Austin", state="TX", postal_code="78701", country="USA",
    )
    session.add(customer)
    session.flush()
    session.add(_FormCustomerRow(
        form_customer="SX5N", form_address="1 Fab Way", form_address_2="Austin, TX",
        customer_id=customer.id,
    ))

    chemical = _ChemicalRow(chemical_name="NH4OH", metals_prep="N/A", silicon_prep="N/A", ions_prep="N/A")
    session.add(chemical)
    session.flush()
    session.add(_FormChemicalRow(form_name="NH4OH", chemical_id=chemical.id))
    # Invented SX5 limits and QC codes, one per site -- the real ones live
    # only in the production database.
    for i, location in enumerate(("SX5N", "SX5S"), start=1):
        session.add(_SpecificationRow(
            id=i, kind="specification", customer_id=customer.id, customer_name_raw=location,
            chemical_id=chemical.id, chemical_name_raw="NH4OH", element_symbol="Al", value=0.1,
            unit="ppb", category_label=f"{location}-QC-TEST", parameter_text="Al", location=location,
        ))

    analysis = _AnalysisRow(analysis_name="36 Elements", analysis_description="36 elements (by ICP-MS)")
    session.add(analysis)
    session.flush()
    session.add(_FormAnalysisRow(form_name="36 Elements", analysis_id=analysis.id))
    session.commit()

    north_folder.mkdir(parents=True, exist_ok=True)
    south_folder.mkdir(parents=True, exist_ok=True)
    services = ReportEngineServices(
        chemical_svc=ChemicalService(session),
        analysis_svc=AnalysisService(session),
        element_svc=ElementService(session),
        customer_svc=CustomerService(session),
        specification_svc=SpecificationService(session),
        report_folder_svc=ReportFolderService(session),
        report_log_svc=ReportGenerationLogService(session),
    )
    services.report_folder_svc.set_folder(customer.id, str(north_folder), created_by="JT", location="SX5N")
    services.report_folder_svc.set_folder(customer.id, str(south_folder), created_by="JT", location="SX5S")
    return services


def _make_sx5_workbook(path: Path) -> None:
    wb = openpyxl.load_workbook(_FORM_PATH)
    ws = wb.active
    ws["C10"] = "SX5N"
    ws["C11"] = "1 Fab Way"
    ws["C12"] = "Austin, TX"
    ws["P5"] = date(2026, 3, 5)
    ws["B22"] = "S-001"
    ws["C22"] = "NH4OH"
    ws["D22"] = "36 Elements"
    ws["K22"] = "Next Day"
    wb.save(path)


def test_sx5_north_and_south_route_to_independent_folders(tmp_path):
    db_url = f"sqlite:///{tmp_path / 'test.db'}"
    north_folder = tmp_path / "Reports" / "SX5_North"
    south_folder = tmp_path / "Reports" / "SX5_South"
    services = _seed_sx5_services(db_url, north_folder, south_folder)

    north_input = tmp_path / "092226-SXOS5 North HCl-1.xlsx"
    south_input = tmp_path / "092226-SXOS5 South HCl-1.xlsx"
    _make_sx5_workbook(north_input)
    _make_sx5_workbook(south_input)

    north_result = run(north_input, None, db_url=None, services=services)
    south_result = run(south_input, None, db_url=None, services=services)

    assert north_result == north_folder / report_filename(north_input.name)
    assert south_result == south_folder / report_filename(south_input.name)
    assert north_result.exists()
    assert south_result.exists()


def test_sx5_ambiguous_filename_raises_before_writing_anything(tmp_path):
    db_url = f"sqlite:///{tmp_path / 'test.db'}"
    services = _seed_sx5_services(
        db_url, tmp_path / "Reports" / "SX5_North", tmp_path / "Reports" / "SX5_South"
    )

    ambiguous_input = tmp_path / "092226-SXOS5 HCl-1.xlsx"  # no "North" or "South"
    _make_sx5_workbook(ambiguous_input)

    with pytest.raises(ReportEngineError, match="North.*South|could not determine"):
        run(ambiguous_input, None, db_url=None, services=services)
