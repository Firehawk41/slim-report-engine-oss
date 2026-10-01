"""End-to-end proof for the SX5 non-routine path. SX5 has no intake form
of its own -- confirmed real (DOMAIN_BRIEF.md): a SX5 non-routine sample
arrives via the SAME Chemical Testing Request form as any other customer,
distinguished only by the resolved customer name. This proves the real
Chemical form parses correctly into a SX5-routable TRSample, and that
sx5_non_routine_report_builder (not chemical_water_report_builder) is the
right dispatch target once the customer resolves to "SX5" -- the single
real customer (North and South are the same customer for invoicing,
confirmed 2026-09-23) -- with North/South itself derived from the TR
filename via submission_report_builder.sx5_location_from_filename, not
from the customer record.
"""

import os
from datetime import date
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from infrastructure.database import Base
from slim_domain.domain.analysis.analysis_service import AnalysisService
from slim_domain.domain.chemical.chemical import Chemical
from slim_domain.domain.chemical.chemical_service import ChemicalService
from slim_domain.domain.customer.customer import Customer
from slim_domain.domain.element.element_repository import _ElementRow
from slim_domain.domain.element.element_service import ElementService
from slim_domain.domain.tr.tr_submission_form_parser import TRSubmissionFormParser

from slim_report_engine.reporting import sx5_non_routine_report_builder, submission_report_builder
from slim_report_engine.reporting.presets import analyte_presets

_FORM_PATH = (
    Path(os.environ.get("PVB_BRIDGE_PATH", r"C:\Users\labuser\Python-VBA-Bridge"))
    / "ReportCreator Project"
    / "Doc-F31 EXAMPLE LAB Analytical Testing Request Form - Chemicals.xlsx"
)

_SX5_CUSTOMER = Customer(
    id=3, name="SX5", street_address="1 Fab Way",
    city="Austin", state="TX", postal_code="78701", country="USA",
)
_NH4OH_CHEMICAL = Chemical(id=9, name="NH4OH", metals_prep="N/A", silicon_prep="N/A", ions_prep="N/A")


class _StubResolver:
    def resolve_customer(self, name, address, address2):
        return _SX5_CUSTOMER

    def resolve_chemical(self, form_name):
        return _NH4OH_CHEMICAL


@pytest.fixture
def session():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as s:
        for symbol in analyte_presets.trace_elements_36():
            s.add(_ElementRow(element_symbol=symbol, element_name=symbol))
        s.commit()
        yield s


@pytest.mark.skipif(not _FORM_PATH.exists(), reason="real Chemical intake form not available in this environment")
def test_real_chemical_form_with_sx5_customer_routes_to_sx5_builder(session):
    wb = openpyxl.load_workbook(_FORM_PATH)
    ws = wb.active

    # SX5 arrives via the exact same form as any other customer -- only
    # the resolved customer/chemical differ (the StubResolver here always
    # resolves to SX5/NH4OH regardless of the literal form text, exactly
    # like a real resolver would once it recognizes these values). The
    # form text itself can still say "SX5N" or similar -- that's just
    # what the technician typed, unrelated to the single real "SX5"
    # customer record it resolves to.
    ws["C10"] = "SX5N"
    ws["C11"] = "1 Fab Way"
    ws["C12"] = "Austin, TX"
    ws["P5"] = date(2026, 3, 5)  # Date received (Chemical column map)

    ws["B22"] = "S-001"
    ws["C22"] = "NH4OH"
    ws["K22"] = "Next Day"  # processing time, column 11 for Chemical

    analysis_svc = AnalysisService(session)
    chemical_svc = ChemicalService(session)
    element_svc = ElementService(session)
    resolver = _StubResolver()

    parser = TRSubmissionFormParser(
        chemical_service=chemical_svc,
        analysis_service=analysis_svc,
        element_service=element_svc,
        input_resolver=resolver,
    )

    submission = parser.build_from_worksheet(ws, filename="092226-SXOS5 North HCl-1.xlsx")

    assert len(submission.samples) == 1
    sample = submission.samples[0]

    customer = _SX5_CUSTOMER  # caller resolves via CustomerService.load_customer(submission.customer_id) in real use
    chemical = chemical_svc.load_chemical(sample.chemical_id) or _NH4OH_CHEMICAL

    assert customer.name == "SX5"  # confirms this sample should route to the SX5 builder, not the generic one

    # location is NOT read off the customer (both North and South are the
    # same "SX5" customer) -- it's derived from the TR filename instead.
    location = submission_report_builder.sx5_location_from_filename(submission.file_name)
    assert location == "SX5N"

    result = sx5_non_routine_report_builder.build_non_routine_report(
        chemical.name, location, submission.date_received, sample.sample_name, element_svc,
        spec_data=sx5_non_routine_report_builder.SX5SpecData(specs={"Al": 0.1}, qc_code="QC-TEST"),
    )

    assert len(result.sections) == 1
    assert result.chemical_label == "NH4OH"
    assert result.name_cell_address == "D5"
    assert "NH4OH" in result.sample_string
    assert "SX5N" in result.sample_string

    data_rows = [r for r in result.sections[0].rows if r.style_name == "DataLabel"]
    assert len(data_rows) == 36
