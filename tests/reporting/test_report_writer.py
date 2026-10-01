import zipfile

import openpyxl
import pytest

from slim_report_engine.reporting import report_writer
from slim_report_engine.reporting.report_row import ReportRow
from slim_report_engine.reporting.report_section import ReportSection


@pytest.fixture
def ws():
    wb = openpyxl.Workbook()
    return wb.active


# ---------------------------------------------------------------------------
# apply_style
# ---------------------------------------------------------------------------

def test_normal_style_leaves_default_formatting(ws):
    report_writer.apply_style(ws, row_index=1, min_col=1, max_col=3, style_name="Normal")
    cell = ws.cell(row=1, column=1)
    assert cell.font.bold is not True
    assert cell.border.left.style is None


def test_section_title_style_is_bold_no_border_no_fill(ws):
    """SX5's chemical-label title row (e.g. "NH4OH") -- confirmed real:
    bold, no border, no fill."""
    report_writer.apply_style(ws, row_index=1, min_col=1, max_col=5, style_name="SectionTitle")
    cell = ws.cell(row=1, column=1)
    assert cell.font.bold is True
    assert cell.border.left.style is None
    assert cell.fill.patternType is None


def test_sample_id_echo_wide_style(ws):
    """Confirmed real (Report Creator Template.xlsx, SX5-N/S, Wafers):
    column 1 = white fill + bold + border; columns 2/3 = pale cyan;
    column 4+ = pale cyan + bold, extending across every Wafer sample
    column. Column 4's value cell is size 12, vertically centered
    (previously size 11, no vertical alignment)."""
    report_writer.apply_style(ws, row_index=1, min_col=1, max_col=6, style_name="SampleIdEchoWide")
    col1 = ws.cell(row=1, column=1)
    assert col1.font.bold is True
    assert col1.fill.fgColor.rgb == "FFFFFFFF"
    assert col1.border.left.style == "thin"
    assert col1.alignment.vertical == "center"

    col2 = ws.cell(row=1, column=2)
    assert col2.fill.fgColor.rgb == "FFCCFFFF"
    assert col2.border.left.style == "thin"
    assert col2.alignment.vertical == "center"

    col4 = ws.cell(row=1, column=4)
    assert col4.fill.fgColor.rgb == "FFCCFFFF"
    assert col4.font.bold is True
    assert col4.font.size == 12
    assert col4.alignment.horizontal == "center"
    assert col4.alignment.vertical == "center"
    assert col4.alignment.wrap_text is True  # confirmed real (D6 in the real template)


def test_sample_id_echo_wide_style_no_border_between_label_and_next_cell():
    """Confirmed real (every "Sample Identification:" row checked): no
    border between the label (column 2) and the cell to its right
    (column 3) -- both sides of that shared edge are omitted."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=1, min_col=1, max_col=6, style_name="SampleIdEchoWide")
    col2 = ws.cell(row=1, column=2)
    col3 = ws.cell(row=1, column=3)
    assert col2.border.right is None or col2.border.right.style is None
    assert col3.border.left is None or col3.border.left.style is None
    # outer edges and top/bottom stay thin
    assert col2.border.left.style == "thin"
    assert col2.border.top.style == "thin"
    assert col2.border.bottom.style == "thin"
    assert col3.border.right.style == "thin"
    assert col3.border.top.style == "thin"
    assert col3.border.bottom.style == "thin"


def test_sample_id_echo_simple_style_has_no_column_1_border():
    """Confirmed real: TOC/Electrical/Misc Analysis's Sample-ID row has NO
    border or fill at all in column 1, unlike the Wide variant."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=1, min_col=1, max_col=4, style_name="SampleIdEchoSimple")
    col1 = ws.cell(row=1, column=1)
    assert col1.border.left.style is None
    assert col1.fill.patternType is None
    col4 = ws.cell(row=1, column=4)
    assert col4.fill.fgColor.rgb == "FFCCFFFF"
    assert col4.font.size == 12
    assert col4.alignment.vertical == "center"


def test_sample_id_echo_simple_style_no_border_between_label_and_next_cell():
    """Same no-border rule as the Wide variant, between columns 2 and 3."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=1, min_col=1, max_col=4, style_name="SampleIdEchoSimple")
    col2 = ws.cell(row=1, column=2)
    col3 = ws.cell(row=1, column=3)
    assert col2.border.right is None or col2.border.right.style is None
    assert col3.border.left is None or col3.border.left.style is None
    assert col2.alignment.vertical == "center"


def test_sample_id_echo_assay_style_only_styles_columns_2_and_3():
    """SX5's Assay-shape "Sample #" row -- confirmed real: shifted one
    column left of the Wide variant, nothing styled outside columns 2/3."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=1, min_col=1, max_col=4, style_name="SampleIdEchoAssay")
    assert ws.cell(row=1, column=1).border.left.style is None
    assert ws.cell(row=1, column=2).fill.fgColor.rgb == "FFCCFFFF"
    assert ws.cell(row=1, column=3).fill.fgColor.rgb == "FFCCFFFF"
    assert ws.cell(row=1, column=4).border.left.style is None


def test_column_header_style_not_bold_white_then_light_blue(ws):
    """Confirmed real: category label (col 2/3) = white, NOT bold;
    Results/Recovery/MDL (col 4+) = light blue, also NOT bold -- despite
    the style being named after modReportStyles.bas's "HeaderBold"."""
    report_writer.apply_style(ws, row_index=2, min_col=1, max_col=6, style_name="ColumnHeader")
    col2 = ws.cell(row=2, column=2)
    assert col2.font.bold is not True
    assert col2.fill.fgColor.rgb == "FFFFFFFF"
    col4 = ws.cell(row=2, column=4)
    assert col4.font.bold is not True
    assert col4.fill.fgColor.rgb == "FF99CCFF"
    assert col4.alignment.horizontal == "center"


def test_column_header_style_labels_vertically_centered():
    """Confirmed real (every "Specification"/"Element"/"Results"/
    "Recovery"/"MDL" label on this row, generic and SX5 both)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=2, min_col=1, max_col=6, style_name="ColumnHeader")
    for col in range(1, 7):
        assert ws.cell(row=2, column=col).alignment.vertical == "center"


def test_column_header_style_column_1_gets_avantgarde_font_no_fill():
    """SX5's element-panel header row has REAL "Specification" text in
    column 1 (AvantGarde, no fill) -- generic panels leave it blank, but
    applying the same font there is harmless since it's invisible."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=2, min_col=1, max_col=6, style_name="ColumnHeader")
    col1 = ws.cell(row=2, column=1)
    assert col1.font.name == "AvantGarde"
    assert col1.fill.patternType is None
    assert col1.border.left.style == "thin"


def test_column_header_assay_style_results_start_at_column_3():
    """SX5's Assay-shape header row -- confirmed real: results (bold,
    light blue) start one column earlier than the generic ColumnHeader
    style, and columns 1/2 have no fill at all."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=5, min_col=1, max_col=4, style_name="ColumnHeaderAssay")
    assert ws.cell(row=5, column=1).fill.patternType is None
    assert ws.cell(row=5, column=2).fill.patternType is None
    col3 = ws.cell(row=5, column=3)
    assert col3.fill.fgColor.rgb == "FF99CCFF"
    assert col3.font.bold is True


def test_data_label_style(ws):
    report_writer.apply_style(ws, row_index=2, min_col=1, max_col=1, style_name="DataLabel")
    cell = ws.cell(row=2, column=1)
    assert cell.border.left.style == "thin"
    assert cell.alignment.horizontal == "center"  # column 1 (spec-value convention)


def test_data_label_style_column_2_plus_is_left_aligned():
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=2, min_col=1, max_col=3, style_name="DataLabel")
    cell = ws.cell(row=2, column=2)
    assert cell.alignment.horizontal == "left"
    assert cell.border.left.style == "thin"


def test_data_value_style(ws):
    report_writer.apply_style(ws, row_index=2, min_col=1, max_col=1, style_name="DataValue")
    cell = ws.cell(row=2, column=1)
    assert cell.border.left.style == "thin"
    assert cell.alignment.horizontal == "right"


def test_summary_row_style_pale_cyan_fill_skips_column_1():
    """The metals panel's AVERAGE/TOTAL rows -- confirmed real: same pale
    cyan fill as the Sample-ID row; no real template populates column 1
    here, so it's left completely unstyled."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=40, min_col=1, max_col=6, style_name="SummaryRow")
    assert ws.cell(row=40, column=1).border.left.style is None
    col2 = ws.cell(row=40, column=2)
    assert col2.fill.fgColor.rgb == "FFCCFFFF"
    assert col2.border.left.style == "thin"


def test_normal_bold_style_no_border_no_fill_but_bold():
    """Wafer's "Notes:" and units-note rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=3, min_col=1, max_col=1, style_name="NormalBold")
    cell = ws.cell(row=3, column=1)
    assert cell.font.bold is True
    assert cell.border.left.style is None
    assert cell.fill.patternType is None


def test_normal_bold_blue_style_bold_and_blue():
    """Confirmed real (Report Creator Template.xlsx row 5): the generic
    preamble's own "Blue font indicates..." line."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=3, min_col=1, max_col=1, style_name="NormalBoldBlue")
    cell = ws.cell(row=3, column=1)
    assert cell.font.bold is True
    assert cell.font.color.rgb == "FF0000FF"
    assert cell.border.left.style is None
    assert cell.fill.patternType is None


def test_normal_blue_style_blue_not_bold():
    """Confirmed real (Wafers Report Template.xlsx row 4): Wafer's own
    equivalent line is blue but NOT bold, unlike the generic preamble's."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=3, min_col=1, max_col=1, style_name="NormalBlue")
    cell = ws.cell(row=3, column=1)
    assert cell.font.bold is not True
    assert cell.font.color.rgb == "FF0000FF"


def test_unknown_style_raises(ws):
    with pytest.raises(ValueError, match="unknown style name"):
        report_writer.apply_style(ws, row_index=1, min_col=1, max_col=1, style_name="Nonexistent")


# ---------------------------------------------------------------------------
# apply_zoom
# ---------------------------------------------------------------------------

def test_apply_zoom_sets_zoom_scale(ws):
    report_writer.apply_zoom(ws, 90)
    assert ws.sheet_view.zoomScale == 90


def test_apply_zoom_wafer_value(ws):
    report_writer.apply_zoom(ws, 95)
    assert ws.sheet_view.zoomScale == 95


# ---------------------------------------------------------------------------
# write_row
# ---------------------------------------------------------------------------

def test_write_row_writes_values(ws):
    row = ReportRow(style_name="Normal")
    row.set_value(2, "Aluminum")
    row.set_value(3, "Al")
    report_writer.write_row(ws, row, row_index=5)
    assert ws.cell(row=5, column=2).value == "Aluminum"
    assert ws.cell(row=5, column=3).value == "Al"


def test_write_row_neutralizes_formula_injection_in_string_values(ws):
    """CWE-1236 / CSV-Excel formula injection: a cell value starting with
    =/+/-/@ is evaluated as a formula when Excel later opens the file --
    guarded here since this writes real customer/technician-typed text."""
    row = ReportRow(style_name="Normal")
    row.set_value(2, "=cmd|'/c calc'!A1")
    row.set_value(3, "+1+1")
    row.set_value(4, "@SUM(1,1)")
    row.set_value(5, "-2+2")
    report_writer.write_row(ws, row, row_index=5)
    assert ws.cell(row=5, column=2).value == "'=cmd|'/c calc'!A1"
    assert ws.cell(row=5, column=3).value == "'+1+1"
    assert ws.cell(row=5, column=4).value == "'@SUM(1,1)"
    assert ws.cell(row=5, column=5).value == "'-2+2"


def test_write_row_leaves_ordinary_values_untouched(ws):
    row = ReportRow(style_name="Normal")
    row.set_value(2, "Aluminum")
    row.set_value(3, -2.5)  # a real negative NUMBER, not a string -- unaffected
    report_writer.write_row(ws, row, row_index=5)
    assert ws.cell(row=5, column=2).value == "Aluminum"
    assert ws.cell(row=5, column=3).value == -2.5


def test_write_row_translates_formula_to_real_a1_range(ws):
    row = ReportRow(style_name="DataLabel")
    row.set_formula(4, "=AVERAGE(R[-4]C:R[-2]C)")
    report_writer.write_row(ws, row, row_index=40)
    assert ws.cell(row=40, column=4).value == "=AVERAGE(D36:D38)"


def test_write_row_applies_style(ws):
    row = ReportRow(style_name="ColumnHeader")
    row.set_value(2, "Element")
    report_writer.write_row(ws, row, row_index=1)
    cell = ws.cell(row=1, column=2)
    assert cell.font.bold is not True
    assert cell.fill.fgColor.rgb == "FFFFFFFF"


def test_write_row_applies_row_height_when_set(ws):
    row = ReportRow(style_name="Normal", row_height=60.0)
    report_writer.write_row(ws, row, row_index=3)
    assert ws.row_dimensions[3].height == 60.0


def test_write_row_leaves_row_height_unset_by_default(ws):
    row = ReportRow(style_name="Normal")
    report_writer.write_row(ws, row, row_index=3)
    assert ws.row_dimensions[3].height is None


def test_write_row_applies_merges(ws):
    row = ReportRow(style_name="Normal")
    row.add_merge(4, 6)
    row.add_merge(2, 3)
    report_writer.write_row(ws, row, row_index=5)
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert "D5:F5" in merged
    assert "B5:C5" in merged


def test_write_row_no_merges_by_default(ws):
    row = ReportRow(style_name="Normal")
    report_writer.write_row(ws, row, row_index=5)
    assert len(ws.merged_cells.ranges) == 0


def test_write_row_applies_real_superscript_rich_text(ws):
    """Confirmed real 2026-09-26 (S:\\Reports\\Customer abf1e): Wafer's own
    "Results in units of atoms/cm2" note has a genuinely superscript
    trailing "2" -- real Excel rich text, not a plain string."""
    from openpyxl.cell.rich_text import CellRichText, TextBlock

    row = ReportRow(style_name="NormalBold")
    row.set_value(1, "Results in units of atoms/cm2")
    row.set_superscript_runs(1, (("Results in units of atoms/cm", False), ("2", True)))
    report_writer.write_row(ws, row, row_index=1)

    cell = ws.cell(row=1, column=1)
    assert isinstance(cell.value, CellRichText)
    assert cell.value[0] == "Results in units of atoms/cm"
    superscript_block = cell.value[1]
    assert isinstance(superscript_block, TextBlock)
    assert superscript_block.text == "2"
    assert superscript_block.font.vertAlign == "superscript"
    # Matches the row's own already-applied font (NormalBold -- Arial,
    # size 10, bold), not some hardcoded font -- confirmed real (the
    # actual Customer abf1e cell's rich text uses this exact base font).
    assert superscript_block.font.rFont == "Arial"
    assert superscript_block.font.sz == 10
    assert superscript_block.font.b is True


def test_write_row_superscript_runs_supports_multiple_segments(ws):
    """The "atoms/cm^2 x 10^10" scaled variant has TWO superscript runs
    (the "10" exponent and the trailing "2"), not just one."""
    row = ReportRow(style_name="NormalBold")
    row.set_value(1, "Results in units of 1010 atoms/cm2")
    row.set_superscript_runs(
        1,
        (
            ("Results in units of 10", False),
            ("10", True),
            (" atoms/cm", False),
            ("2", True),
        ),
    )
    report_writer.write_row(ws, row, row_index=1)

    cell = ws.cell(row=1, column=1)
    superscript_texts = [b.text for b in cell.value if hasattr(b, "font") and b.font.vertAlign == "superscript"]
    assert superscript_texts == ["10", "2"]
    plain_texts = [b for b in cell.value if isinstance(b, str)]
    assert plain_texts == ["Results in units of 10", " atoms/cm"]


def test_write_row_no_superscript_runs_by_default(ws):
    row = ReportRow(style_name="Normal")
    row.set_value(1, "plain text")
    report_writer.write_row(ws, row, row_index=5)
    assert ws.cell(row=5, column=1).value == "plain text"


def test_write_row_note_marker_is_bold_red_with_no_border(ws):
    """Confirmed real (2026-09-28, a real Folder 3978d Z-report): the
    KED-mode "Note 1" marker sits one column past the panel's own
    max_columns=6 bordered box, bold red, with NO border/fill of its own."""
    row = ReportRow(style_name="DataLabel", max_columns=6)
    row.set_value(2, "Nickel")
    row.set_value(3, "Ni")
    row.set_note(7, "Note 1")
    report_writer.write_row(ws, row, row_index=5)

    note_cell = ws.cell(row=5, column=7)
    assert note_cell.value == "Note 1"
    assert note_cell.font.bold is True
    assert note_cell.font.color.rgb == "FFFF0000"
    assert note_cell.border.left is None or note_cell.border.left.style is None
    assert note_cell.border.right is None or note_cell.border.right.style is None

    # The symbol cell itself stays plain "DataLabel" styling -- NOT bold
    # red (the 2026-09-25 behavior this replaces).
    symbol_cell = ws.cell(row=5, column=3)
    assert symbol_cell.font.bold is not True
    assert symbol_cell.font.color is None


def test_write_row_unsupported_formula_shape_raises(ws):
    row = ReportRow(style_name="Normal")
    row.set_formula(1, "=SUM(R1C1:R5C5)")  # a cross-column shape, never used, not supported
    with pytest.raises(ValueError, match="doesn't match either supported"):
        report_writer.write_row(ws, row, row_index=1)


def test_write_row_translates_same_column_list_formula_to_real_a1_cells(ws):
    """Hardness's shape -- confirmed real (modReportCreator.bas
    Set_Hardness_Formula/Formula_To_Sum_Elements): specific, non-
    contiguous element rows summed by a live formula, not a range."""
    row = ReportRow(style_name="DataLabel")
    row.set_formula(4, "=SUM(R[-6]C,R[-5]C,R[-2]C)")
    report_writer.write_row(ws, row, row_index=40)
    assert ws.cell(row=40, column=4).value == "=SUM(D34,D35,D38)"


def test_write_row_translates_single_item_list_formula(ws):
    row = ReportRow(style_name="DataLabel")
    row.set_formula(4, "=SUM(R[-5]C)")
    report_writer.write_row(ws, row, row_index=9)
    assert ws.cell(row=9, column=4).value == "=SUM(D4)"


# ---------------------------------------------------------------------------
# write_section / write_sections
# ---------------------------------------------------------------------------

def _two_row_section(id: str) -> ReportSection:
    section = ReportSection(id)
    r1 = ReportRow(style_name="Normal")
    r1.set_value(1, f"{id}-row1")
    section.add_row(r1)
    r2 = ReportRow(style_name="Normal")
    r2.set_value(1, f"{id}-row2")
    section.add_row(r2)
    return section


def test_write_section_returns_next_row(ws):
    next_row = report_writer.write_section(ws, _two_row_section("A"), start_row=1)
    assert next_row == 3
    assert ws.cell(row=1, column=1).value == "A-row1"
    assert ws.cell(row=2, column=1).value == "A-row2"


def test_write_sections_chains_back_to_back(ws):
    sections = [_two_row_section("A"), _two_row_section("B")]
    next_row = report_writer.write_sections(ws, sections, start_row=1)
    assert next_row == 5
    assert ws.cell(row=1, column=1).value == "A-row1"
    assert ws.cell(row=2, column=1).value == "A-row2"
    assert ws.cell(row=3, column=1).value == "B-row1"
    assert ws.cell(row=4, column=1).value == "B-row2"


# ---------------------------------------------------------------------------
# apply_header_footer
# ---------------------------------------------------------------------------

def test_apply_header_footer_sets_expected_fields(ws):
    report_writer.apply_header_footer(
        ws,
        center_header="36 Elements",
        customer_name="Customer ea6f9",
        customer_address_line1="123 Main St",
        customer_address_line2="Springfield, IL",
        lab_header_lines=("Some Lab", "100 Industrial Pkwy", "Phone: 555-0100"),
    )
    assert ws.oddHeader.left.text == "Some Lab\n100 Industrial Pkwy\nPhone: 555-0100"
    assert ws.oddHeader.center.text == "36 Elements"
    assert "Customer ea6f9" in ws.oddHeader.right.text
    assert ws.oddFooter.center.text == "&F"
    assert ws.oddFooter.right.text == "Page &P"
    assert ws.page_margins.top == 0.99


def test_apply_header_footer_survives_a_real_save_and_reload(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_header_footer(
        ws,
        center_header="36 Elements",
        customer_name="Customer ea6f9",
        customer_address_line1="123 Main St",
        customer_address_line2="Springfield, IL",
    )
    path = tmp_path / "test.xlsx"
    wb.save(path)

    reloaded = openpyxl.load_workbook(path)
    reloaded_ws = reloaded.active
    assert reloaded_ws.oddHeader.center.text == "36 Elements"
    assert reloaded_ws.oddFooter.center.text == "&F"
    assert reloaded_ws.page_margins.top == 0.99


def test_apply_header_footer_escapes_ampersands_for_excel(tmp_path):
    """Real customer and chemical names contain "&" (found in the live
    database). A bare "&" is an Excel header formatting code, so
    it must reach the saved file as "&&" -- while the footer's own
    deliberate &F/&P codes stay single."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_header_footer(
        ws,
        center_header="Water & Sulfolane",
        customer_name="Smith & Sons Engineering",
        customer_address_line1="1 A&B Way",
        customer_address_line2="Dallas, TX",
        lab_header_lines=("Lab & Co",),
    )
    path = tmp_path / "test.xlsx"
    wb.save(path)

    sheet_xml = zipfile.ZipFile(path).read("xl/worksheets/sheet1.xml").decode()
    # The XML itself escapes each "&" as "&amp;", so "&&" appears doubled.
    assert "Water &amp;&amp; Sulfolane" in sheet_xml
    assert "Smith &amp;&amp; Sons Engineering" in sheet_xml
    assert "1 A&amp;&amp;B Way" in sheet_xml
    assert "Lab &amp;&amp; Co" in sheet_xml
    assert "&amp;F" in sheet_xml and "Page &amp;P" in sheet_xml
    assert "&amp;&amp;F" not in sheet_xml and "&amp;&amp;P" not in sheet_xml


def test_apply_header_footer_no_lab_lines_leaves_left_header_unset(ws):
    """Deliberately left unset (not assigned an empty string) -- see the
    comment in apply_header_footer for the real openpyxl serialization bug
    this avoids."""
    report_writer.apply_header_footer(
        ws,
        center_header="TOC",
        customer_name="Customer ea6f9",
        customer_address_line1="",
        customer_address_line2="",
    )
    assert ws.oddHeader.left.text is None


def test_apply_header_footer_no_lab_lines_survives_save_and_reload(tmp_path):
    """Regression test for the empty-string-corrupts-neighboring-parts bug:
    confirms center/right still come back correctly when left is never set."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_header_footer(
        ws,
        center_header="TOC",
        customer_name="Customer ea6f9",
        customer_address_line1="",
        customer_address_line2="",
    )
    path = tmp_path / "test.xlsx"
    wb.save(path)

    reloaded = openpyxl.load_workbook(path)
    reloaded_ws = reloaded.active
    assert reloaded_ws.oddHeader.center.text == "TOC"
    assert reloaded_ws.oddHeader.right.text is not None
    assert "Customer ea6f9" in reloaded_ws.oddHeader.right.text


# ---------------------------------------------------------------------------
# apply_column_widths
# ---------------------------------------------------------------------------

def test_apply_column_widths_sets_each_listed_column(ws):
    report_writer.apply_column_widths(ws, {1: 17.14, 3: 6.57})
    assert ws.column_dimensions["A"].width == 17.14
    assert ws.column_dimensions["C"].width == 6.57


def test_apply_column_widths_leaves_unlisted_columns_alone(ws):
    report_writer.apply_column_widths(ws, {1: 17.14})
    assert "B" not in ws.column_dimensions


# ---------------------------------------------------------------------------
# apply_processing_time_stamp
# ---------------------------------------------------------------------------

def test_apply_processing_time_stamp_red(ws):
    report_writer.apply_processing_time_stamp(ws, "D1", "Same Day RUSH", "red", 10)
    cell = ws["D1"]
    assert cell.value == "Same Day RUSH"
    assert cell.font.bold is True
    assert cell.font.size == 10
    assert cell.font.color.rgb == "FFFF0000"


def test_apply_processing_time_stamp_blue(ws):
    report_writer.apply_processing_time_stamp(ws, "E1", "Next Day Time Limited", "blue", 12)
    cell = ws["E1"]
    assert cell.value == "Next Day Time Limited"
    assert cell.font.size == 12
    assert cell.font.color.rgb == "FF0000FF"


def test_apply_plain_cell_stamp_clears_a_preexisting_bold_font(ws):
    """2026-09-29, SX5's first-ever real golden-diff run: SX5's QC row
    (SampleIdEchoWide style) already leaves this exact cell bold before
    the external stamp ever runs -- confirmed real legacy output shows
    the sample-ID string plain, not bold. Reproduces that pre-existing
    state directly rather than trusting a fresh cell's default (which
    would pass even without the fix)."""
    from openpyxl.styles import Font

    ws["D5"].font = Font(name="Arial", size=12, bold=True)
    report_writer.apply_plain_cell_stamp(ws, "D5", "092926-49%HF-SX5S-49%HF-2")
    cell = ws["D5"]
    assert cell.value == "092926-49%HF-SX5S-49%HF-2"
    assert cell.font.bold is not True


def test_apply_plain_cell_stamp_on_a_fresh_cell(ws):
    report_writer.apply_plain_cell_stamp(ws, "B2", "some value")
    cell = ws["B2"]
    assert cell.value == "some value"
    assert cell.font.bold is not True


# ---------------------------------------------------------------------------
# limit columns (warning limits shift content right by one)
# ---------------------------------------------------------------------------

_TWO_LIMITS = (("specification", "Action Limit"), ("warning_limit", "Warning Limit"))


def test_single_limit_column_layout_is_unchanged(ws):
    row = ReportRow(style_name="DataLabel")
    row.set_value(2, "Aluminum")
    row.set_value(3, "Al")
    row.set_limit("specification", 0.4)
    report_writer.write_row(ws, row, 1, (("specification", "Specification"),))
    assert [ws.cell(row=1, column=c).value for c in (1, 2, 3)] == [0.4, "Aluminum", "Al"]


def test_two_limit_columns_place_both_limits_and_shift_content(ws):
    row = ReportRow(style_name="DataLabel")
    row.set_value(2, "Aluminum")
    row.set_value(3, "Al")
    row.set_limit("specification", 0.4)
    row.set_limit("warning_limit", 0.3)
    report_writer.write_row(ws, row, 1, _TWO_LIMITS)
    assert [ws.cell(row=1, column=c).value for c in (1, 2, 3, 4)] == [0.4, 0.3, "Aluminum", "Al"]
    # the extra limit column takes column 1's formatting
    a1, b1 = ws.cell(row=1, column=1), ws.cell(row=1, column=2)
    assert (b1.font.name, b1.font.b, b1.font.sz) == (a1.font.name, a1.font.b, a1.font.sz)
    assert (b1.border.left.style, b1.border.top.style) == (a1.border.left.style, a1.border.top.style) == ("thin", "thin")


def test_limit_header_row_gets_every_label_and_column_1_text_stays_put(ws):
    header = ReportRow(style_name="ColumnHeader")
    header.is_limit_header = True
    header.set_value(2, "Element")
    header.set_value(4, "EXAMPLE LAB Results")
    report_writer.write_row(ws, header, 1, _TWO_LIMITS)
    assert [ws.cell(row=1, column=c).value for c in (1, 2, 3, 5)] == [
        "Action Limit", "Warning Limit", "Element", "EXAMPLE LAB Results"]

    notes = ReportRow()
    notes.set_value(1, "Notes:")
    report_writer.write_row(ws, notes, 2, _TWO_LIMITS)
    assert ws.cell(row=2, column=1).value == "Notes:"


def test_formulas_merges_and_notes_shift_with_their_column(ws):
    row = ReportRow(style_name="SummaryRow")
    row.set_formula(4, "=AVERAGE(R[-3]C:R[-1]C)")
    row.add_merge(4, 6)
    row.set_note(7, "Note 1")
    report_writer.write_row(ws, row, 10, _TWO_LIMITS)
    assert ws.cell(row=10, column=5).value == "=AVERAGE(E7:E9)"
    assert "E10:G10" in [str(r) for r in ws.merged_cells.ranges]
    assert ws.cell(row=10, column=8).value == "Note 1"


def test_stamp_addresses_and_widths_shift(ws):
    assert report_writer.shift_cell_address("D1", _TWO_LIMITS) == "E1"
    assert report_writer.shift_cell_address("A1", _TWO_LIMITS) == "A1"
    assert report_writer.shift_cell_address("D1") == "D1"
    report_writer.apply_column_widths(ws, {1: 17.0, 2: 16.0}, _TWO_LIMITS)
    assert (ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width) == (17.0, 17.0, 16.0)
