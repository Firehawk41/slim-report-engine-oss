"""Real-evidence diff engine for golden-file report tests.

Generalizes the ad hoc diff scripts run by hand throughout the
2026-09-28 session into a reusable, importable tool, so a real content
or styling bug gets caught by `pytest -q` instead of a human visually
reviewing a generated workbook. See the `report-engine-validation-
evidence` project memory for the methodology and the specific real bugs
this formalization is a direct response to (a KED-mode note styled on
the wrong cell, worded wrong, and ordered wrong -- three distinct real
bugs a plain "does this text exist somewhere" check would have missed).

Compares a freshly-generated report workbook against a real legacy-
macro-generated workbook for the SAME real Testing Request:
  - every non-exempt string VALUE in NEW must appear somewhere on the
    same-named sheet in the legacy output (catches missing/wrong text,
    missing dispatch cases, wrong labels);
  - every NEW cell with notably bold and/or colored font must have a
    matching-styled instance of that same text in the legacy output
    (catches styling bugs like a marker painted on the wrong cell);
  - within each block of rows (sections are always separated by exactly
    one blank row -- see row_builders.add_footer_row), the relative
    order of "Analysis by ...", "Test Methods: ...", and "Note ..."
    marker lines in NEW must match their order in the matching legacy
    block (catches ordering bugs).

Deliberately NOT a strict cell-by-cell positional diff: real legacy
reports vary in ways that are known-good, not bugs (see _is_exempt).
"""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

SKIP_SHEETS = frozenset({"Analytical Testing Request"})

# Deliberately one-sided and permanent, not a bug (direct user
# instruction, 2026-09-28): SOP/"Test Methods:" lines are a genuine
# Python-side upgrade over the real legacy macro, which never generated
# them at all -- diffing them against legacy output would be comparing
# against a feature that was never supposed to exist there.
_EXEMPT_PREFIXES = ("Test Methods:",)

# GC-FID's chemical-name placeholders are hand-filled by a technician
# after generation in the real completed report -- never present
# verbatim in a freshly generated NEW workbook.
_EXEMPT_VALUES = frozenset({"Chemical 1", "Chemical 2"})

_MARKER_PREFIXES = ("Analysis by", "Test Methods:", "Note ")


@dataclass(frozen=True)
class CellStyle:
    bold: bool
    # A coarse boolean, not an exact color -- real workbooks mix THREE
    # different Excel color representations for the same perceived color
    # (literal RGB, the legacy indexed palette, and theme+tint), and
    # resolving indexed/theme to an exact RGB needs the workbook's own
    # theme XML. Confirmed real (2026-09-28): the real legacy macro's own
    # "Blue font indicates..." line uses indexed=12, not an RGB -- an
    # exact-RGB comparison against this project's own literal-RGB blue
    # flagged it as a false-positive mismatch. Comparing "is this cell
    # notably colored at all" catches the real bug class (a marker
    # painted on the wrong cell, or missing where it should be) without
    # that noise.
    notably_colored: bool
    color_repr: str  # human-readable, for failure messages only -- not compared


@dataclass(frozen=True)
class Mismatch:
    sheet: str
    coordinate: str
    value: object
    kind: str  # "missing_sheet" | "missing_value" | "style_mismatch" | "order_mismatch"
    detail: str

    def __str__(self) -> str:
        return f"[{self.sheet}!{self.coordinate}] {self.kind}: {self.detail} (value={self.value!r})"


_BLACK_RGB = ("FF000000", "00000000")
# Excel's legacy indexed palette: 64 = "automatic" (no explicit color),
# 8 = plain "Black" -- confirmed real (2026-09-28, a real wafer dummy-TR
# diff): the real legacy macro's own "Notes:"/units-note lines use
# indexed=8, which this tool was flagging as a false-positive style
# mismatch against this project's un-colored (default black) text.
_AUTOMATIC_INDEXED = (64, 8)


def _style_of(cell) -> CellStyle:
    font = cell.font
    if font is None:
        return CellStyle(bold=False, notably_colored=False, color_repr="none")

    color = font.color
    notably_colored = False
    color_repr = "none"
    if color is not None:
        # A real openpyxl quirk (2026-09-28, this project's own first
        # encounter): a Color object only genuinely sets the ONE field
        # matching its own .type ("rgb"/"indexed"/"theme") -- reading any
        # OTHER field back doesn't return None or raise, it returns a
        # bogus validation-error placeholder object (e.g. reading .theme
        # on an indexed-type color gives something whose str() is
        # "Values must be of type <class 'int'>", not a real theme id).
        # `is not None` treated that placeholder as a real value; these
        # explicit isinstance checks don't.
        rgb = getattr(color, "rgb", None)
        indexed = getattr(color, "indexed", None)
        theme = getattr(color, "theme", None)
        if isinstance(rgb, str) and rgb not in _BLACK_RGB:
            notably_colored = True
            color_repr = f"rgb={rgb}"
        elif isinstance(indexed, int) and indexed not in _AUTOMATIC_INDEXED:
            notably_colored = True
            color_repr = f"indexed={indexed}"
        elif isinstance(theme, int):
            notably_colored = True
            color_repr = f"theme={theme}"

    return CellStyle(bold=bool(font.bold), notably_colored=notably_colored, color_repr=color_repr)


def _is_exempt(value: object) -> bool:
    if not isinstance(value, str):
        return True
    text = value.strip()
    if not text:
        return True
    if text in _EXEMPT_VALUES:
        return True
    return any(text.startswith(p) for p in _EXEMPT_PREFIXES)


def _sheet_value_index(ws: Worksheet) -> dict[str, list[tuple[str, CellStyle]]]:
    """Maps a cell's stripped string value -> every (coordinate, style)
    it's found at on this sheet (usually one hit, sometimes several for
    common labels like "STDEV")."""
    index: dict[str, list[tuple[str, CellStyle]]] = {}
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.strip():
                index.setdefault(v.strip(), []).append((cell.coordinate, _style_of(cell)))
    return index


def _row_blocks(ws: Worksheet) -> list[list[int]]:
    """Splits a sheet's rows into blocks (lists of row numbers with any
    content) separated by fully-blank rows -- mirrors row_builders'
    own "exactly one blank row separates sections" convention."""
    blocks: list[list[int]] = []
    current: list[int] = []
    for r in range(1, ws.max_row + 1):
        has_content = any(
            ws.cell(row=r, column=c).value not in (None, "") for c in range(1, ws.max_column + 1)
        )
        if has_content:
            current.append(r)
        elif current:
            blocks.append(current)
            current = []
    if current:
        blocks.append(current)
    return blocks


def _block_marker_sequence(ws: Worksheet, rows: list[int]) -> list[str]:
    seq: list[str] = []
    for r in rows:
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str):
                text = v.strip()
                for prefix in _MARKER_PREFIXES:
                    if text.startswith(prefix):
                        seq.append(prefix)
                        break
    return seq


def _block_anchor_text(ws: Worksheet, rows: list[int]) -> str | None:
    """A representative text used to match this block to its counterpart
    on the other sheet -- the "Analysis by ..." footer line, since it's
    normally unique per real panel type within one sheet."""
    for r in rows:
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.strip().startswith("Analysis by"):
                return v.strip()
    return None


def _is_number(v: object) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _limit_index(
    ws: Worksheet, formats_out: dict[str, list[str]] | None = None
) -> dict[str, list[tuple[str, tuple[tuple[int, float], ...]]]]:
    """symbol -> [(symbol's column letter, ((limit column, limit), ...))]
    for every analyte row carrying a limit. A row reads left to right as
    limit slot(s) -- numbers or blanks -- then the analyte name, then its
    symbol: A=limit, B=name, C=symbol normally; A=action, B=warning,
    C=name, D=symbol on a sheet with warning limits (legacy inserted a
    column B). Recording the symbol's column also proves that shift. The
    string check exempts every number (real legacy reports also hold
    hand-entered results), so limits need this targeted check."""
    index: dict[str, list[tuple[str, tuple[tuple[int, float], ...]]]] = {}
    for row in ws.iter_rows(max_col=6):
        limits: list[tuple[int, float]] = []
        formats: list[str] = []
        texts: list[tuple[str, str]] = []
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.strip():
                texts.append((cell.column_letter, v.strip()))
                if len(texts) == 2:
                    break
            elif _is_number(v):
                if texts:
                    break  # a number after the name is a result, not a limit
                limits.append((cell.column, float(v)))
                formats.append(cell.number_format)
        if limits and len(texts) == 2:
            symbol_col, symbol = texts[1]
            index.setdefault(symbol, []).append((symbol_col, tuple(limits)))
            if formats_out is not None:
                formats_out.setdefault(symbol, []).extend(formats)
    return index


def _shows_3dp(number_format: str) -> bool:
    return number_format.endswith(".000")


def _limit_mismatches(sheet_name: str, ws_new: Worksheet, ws_legacy: Worksheet) -> list[Mismatch]:
    """Compares limits symbol by symbol: same values in the same limit
    columns, and the symbol in the same column. The coordinate is reported
    as "limit:<symbol>" so a tolerated exception in KNOWN_OPEN_MISMATCHES
    survives rows moving."""
    new_formats: dict[str, list[str]] = {}
    legacy_formats: dict[str, list[str]] = {}
    new, legacy = _limit_index(ws_new, new_formats), _limit_index(ws_legacy, legacy_formats)
    found: list[Mismatch] = []
    # Display format: wherever legacy shows a limit to 3 decimals (as every
    # real report does), NEW must too -- the value stays a number either way.
    for symbol in sorted(set(new_formats) & set(legacy_formats)):
        if all(_shows_3dp(f) for f in legacy_formats[symbol]) and not all(_shows_3dp(f) for f in new_formats[symbol]):
            found.append(Mismatch(sheet_name, f"limit_format:{symbol}", str(new_formats[symbol]), "limit_format",
                                  f"legacy shows {symbol}'s limit to 3 decimals, NEW uses {new_formats[symbol]}"))
    for symbol in sorted(set(new) | set(legacy)):
        new_vals = sorted(new.get(symbol, []))
        legacy_vals = sorted(legacy.get(symbol, []))
        key = f"limit:{symbol}"
        if not legacy_vals:
            found.append(Mismatch(sheet_name, key, str(new_vals), "extra_limit",
                                  f"NEW prints limit(s) {new_vals}, legacy has none for {symbol}"))
        elif not new_vals:
            found.append(Mismatch(sheet_name, key, str(legacy_vals), "missing_limit",
                                  f"legacy has limit(s) {legacy_vals} for {symbol}, NEW prints none"))
        elif not _same_limits(new_vals, legacy_vals):
            found.append(Mismatch(sheet_name, key, str(new_vals), "limit_mismatch",
                                  f"NEW {new_vals} != legacy {legacy_vals} for {symbol}"))
    return found


def _same_limits(a: list, b: list) -> bool:
    if len(a) != len(b):
        return False
    for (col_a, lims_a), (col_b, lims_b) in zip(a, b):
        if col_a != col_b or len(lims_a) != len(lims_b):
            return False
        for (ca, va), (cb, vb) in zip(lims_a, lims_b):
            if ca != cb or abs(va - vb) > 1e-9 * max(1.0, abs(vb)):
                return False
    return True


def _is_boxed(cell) -> bool:
    """True if the cell is drawn as part of a table: any border or fill."""
    b = cell.border
    bordered = b is not None and any(
        side is not None and side.style for side in (b.left, b.right, b.top, b.bottom)
    )
    return bordered or bool(cell.fill is not None and cell.fill.fill_type)


def diff_workbooks(
    new_path: str,
    legacy_path: str,
    customer_name_substitute: str | None = None,
    real_customer_name: str | None = None,
) -> list[Mismatch]:
    """Diffs new_path (this project's freshly generated output) against
    legacy_path (the real legacy macro's own output for the SAME real
    Testing Request), sheet by sheet. Returns every mismatch found;
    empty means clean.

    customer_name_substitute / real_customer_name: when NEW had to be
    generated with a substituted customer (e.g. "EXAMPLE LAB" itself isn't
    a resolvable customer row -- see _ForgivingCustomerService in
    scripts/scan_latest_request.py), pass the substitute name used and
    the real one it stands in for, so sample-ID echo text differing only
    by that substitution isn't reported as a false-positive mismatch.
    """
    mismatches: list[Mismatch] = []
    wb_new = load_workbook(new_path, data_only=True)
    wb_legacy = load_workbook(legacy_path, data_only=True)

    for sheet_name in wb_new.sheetnames:
        if sheet_name in SKIP_SHEETS:
            continue
        if sheet_name not in wb_legacy.sheetnames:
            mismatches.append(Mismatch(sheet_name, "-", None, "missing_sheet", "sheet absent from legacy output"))
            continue

        ws_new = wb_new[sheet_name]
        ws_legacy = wb_legacy[sheet_name]
        legacy_index = _sheet_value_index(ws_legacy)
        mismatches.extend(_limit_mismatches(sheet_name, ws_new, ws_legacy))

        for row in ws_new.iter_rows():
            for cell in row:
                value = cell.value
                if _is_exempt(value):
                    continue
                text = value.strip()

                compare_text = text
                if customer_name_substitute and real_customer_name and customer_name_substitute in text:
                    compare_text = text.replace(customer_name_substitute, real_customer_name)

                hits = legacy_index.get(compare_text) or legacy_index.get(text)
                if not hits:
                    mismatches.append(
                        Mismatch(sheet_name, cell.coordinate, text, "missing_value", "not found anywhere in legacy sheet")
                    )
                    continue

                new_style = _style_of(cell)
                # The reverse direction (added 2026-10-02 -- the SX5 "Data =<
                # MDL" legend lost its blue and nothing noticed): text legacy
                # ALWAYS colors must not come out uncolored. Color only:
                # color carries meaning (blue = at or below MDL, red = out of
                # spec), while legacy's bold varies by template.
                if not new_style.notably_colored and all(style.notably_colored for _coord, style in hits):
                    mismatches.append(
                        Mismatch(
                            sheet_name, cell.coordinate, text, "style_mismatch",
                            f"NEW is uncolored, legacy always colors this text "
                            f"(legacy styles found: {[(s.bold, s.color_repr) for _c, s in hits]})",
                        )
                    )
                if new_style.bold or new_style.notably_colored:
                    if not any(
                        style.bold == new_style.bold and style.notably_colored == new_style.notably_colored
                        for _coord, style in hits
                    ):
                        mismatches.append(
                            Mismatch(
                                sheet_name, cell.coordinate, text, "style_mismatch",
                                f"NEW is bold={new_style.bold} color=({new_style.color_repr}), legacy has no "
                                f"matching styled instance of this text (legacy styles found: "
                                f"{[(s.bold, s.color_repr) for _c, s in hits]})",
                            )
                        )

        new_blocks = _row_blocks(ws_new)
        legacy_blocks_by_anchor: dict[str, list[list[int]]] = {}
        for rows in _row_blocks(ws_legacy):
            anchor = _block_anchor_text(ws_legacy, rows)
            if anchor:
                legacy_blocks_by_anchor.setdefault(anchor, []).append(rows)

        for rows in new_blocks:
            anchor = _block_anchor_text(ws_new, rows)
            if not anchor or anchor not in legacy_blocks_by_anchor:
                continue
            new_seq = _block_marker_sequence(ws_new, rows)
            if "Note " not in new_seq and not any(m.startswith("Note ") for m in new_seq):
                continue  # nothing to order-check in this block
            legacy_rows = legacy_blocks_by_anchor[anchor][0]
            legacy_seq = _block_marker_sequence(ws_legacy, legacy_rows)
            if new_seq != legacy_seq and set(new_seq) == set(legacy_seq):
                mismatches.append(
                    Mismatch(
                        sheet_name, f"rows {rows[0]}-{rows[-1]}", anchor, "order_mismatch",
                        f"NEW marker order {new_seq} != legacy order {legacy_seq} (same markers, different order)",
                    )
                )

    return mismatches
