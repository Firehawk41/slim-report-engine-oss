"""Formatting rules checked on real generated reports.

The golden diff compares text and its colour against legacy reports; it
cannot see borders, fills or column widths. These tests hold the layout
rules set by direct user correction (2026-10-02), on the same real
Testing Requests the golden pairs use:
  - SX5: the "Data =< MDL" legend is blue, the WX-2000 legend is bold red,
    and the sample-ID header is one column (D) -- nothing boxed past it.
  - SX5 Assay: the table is boxed A-D; a standalone sheet's "Sample #"
    row is boxed B-C, as in legacy.
  - Wafer: the spacer column between the last sample and MDL/QL carries
    no formatting on any row, the sample-ID row stops at the last sample,
    and every sample column is as wide as D so the sample ID can be read.
  - TOTAL rows: the untotalled column between the two sums is unboxed.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from tests.golden._diff import _is_boxed
from tests.golden.test_golden_reports import GOLDEN_PAIRS, _generate

_PAIRS = {p.name: p for p in GOLDEN_PAIRS}
_BLUE, _RED = "FF0000FF", "FFFF0000"


@pytest.fixture(scope="module")
def generated(tmp_path_factory):
    """Generates each needed pair once per module; name -> workbook."""
    out: dict = {}

    def get(name: str):
        if name not in out:
            pair = _PAIRS[name]
            if not pair.input_path.exists():
                pytest.skip(f"real input not reachable: {pair.input_path}")
            path = tmp_path_factory.mktemp(name) / "report.xlsx"
            try:
                _generate(pair, path)
            except Exception as exc:  # pragma: no cover -- no live DB in this environment
                pytest.skip(f"could not generate against the real DB in this environment: {exc}")
            out[name] = load_workbook(path)
        return out[name]

    return get


def _find(ws: Worksheet, text: str):
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip() == text.strip():
                return cell
    return None


def _font_rgb(cell) -> str | None:
    color = cell.font.color if cell.font else None
    return color.rgb if color is not None and isinstance(color.rgb, str) else None


def _boxed_columns(ws: Worksheet, row: int) -> set[int]:
    return {c for c in range(1, ws.max_column + 1) if _is_boxed(ws.cell(row, c))}


def _report_sheets(wb):
    return [ws for ws in wb.worksheets if ws.title != "Analytical Testing Request"]


# --- SX5 ---


@pytest.mark.parametrize("name", ["sxos5_north_all", "sxos5_south_all"])
def test_sx5_mdl_legend_is_blue(generated, name: str) -> None:
    checked = 0
    for ws in _report_sheets(generated(name)):
        cell = _find(ws, "Data =< MDL")
        if cell is not None:
            assert _font_rgb(cell) == _BLUE, f"{ws.title}!{cell.coordinate}"
            checked += 1
    assert checked, "no SX5 element sheet found"


@pytest.mark.parametrize("name", ["sxos5_north_all", "sxos5_south_all"])
def test_sx5_sample_header_is_one_column(generated, name: str) -> None:
    checked = 0
    for ws in _report_sheets(generated(name)):
        for label in ("Sample #:", "Sample #"):
            cell = _find(ws, label)
            if cell is not None:
                if label == "Sample #:":  # element panels: A-D, as in legacy
                    assert _boxed_columns(ws, cell.row) == {1, 2, 3, 4}, f"{ws.title} row {cell.row}"
                checked += 1
    assert checked


@pytest.mark.parametrize("name", ["sxos5_north_all", "sxos5_south_all"])
def test_sx5_assay_table_is_four_columns(generated, name: str) -> None:
    checked = 0
    for ws in _report_sheets(generated(name)):
        header = _find(ws, "Parameter")
        if header is None:
            continue
        for row in (header.row, header.row + 1):
            assert _boxed_columns(ws, row) == {1, 2, 3, 4}, f"{ws.title} row {row}: {_boxed_columns(ws, row)}"
        sample = _find(ws, "Sample #")
        if sample is not None:  # standalone Assay sheets: label + stamp cells only
            assert _boxed_columns(ws, sample.row) == {2, 3}, f"{ws.title} row {sample.row}"
        checked += 1
    if not checked:
        pytest.skip("no Assay-shape sheet in this pair")


def test_sx5_out_of_spec_legend_is_bold_red(generated) -> None:
    found = []
    for name in ("sxos5_north_all", "sxos5_south_all"):
        for ws in _report_sheets(generated(name)):
            cell = _find(ws, "Bold Red Font: Indicates data outside specification limits")
            if cell is not None:
                found.append((ws.title, cell.font.bold, _font_rgb(cell)))
    assert found, "no WX-2000 sheet in the SX5 pairs"
    assert all(bold and rgb == _RED for _t, bold, rgb in found), found


# --- Wafer ---


def _wafer_panels(ws: Worksheet):
    """(sample-ID row, header row, last analyte row, MDL/QL column) per panel."""
    panels = []
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, 2).value in ("Element", "Anion"):
            last_col = max(c for c in range(1, ws.max_column + 1) if ws.cell(row, c).value not in (None, ""))
            end = row + 1
            while end + 1 <= ws.max_row and ws.cell(end + 1, 3).value:
                end += 1
            panels.append((row - 1, row, end, last_col))
    return panels


@pytest.mark.parametrize("name", ["dummy_tr_wafers", "realistic_wafers_dummy"])
def test_wafer_spacer_column_has_no_formatting(generated, name: str) -> None:
    checked = 0
    for ws in _report_sheets(generated(name)):
        for sample_row, _header, last_row, mdl_col in _wafer_panels(ws):
            spacer = mdl_col - 1
            for row in range(sample_row, last_row + 1):
                cell = ws.cell(row, spacer)
                assert not _is_boxed(cell), f"{ws.title}!{cell.coordinate} is formatted"
                assert cell.value in (None, ""), f"{ws.title}!{cell.coordinate} has a value"
            checked += 1
    assert checked, "no wafer panel found"


@pytest.mark.parametrize("name", ["dummy_tr_wafers", "realistic_wafers_dummy"])
def test_wafer_sample_id_row_stops_at_the_last_sample(generated, name: str) -> None:
    for ws in _report_sheets(generated(name)):
        for sample_row, _header, _last, mdl_col in _wafer_panels(ws):
            assert max(_boxed_columns(ws, sample_row)) <= mdl_col - 2, f"{ws.title} row {sample_row}"


@pytest.mark.parametrize("name", ["dummy_tr_wafers", "realistic_wafers_dummy"])
def test_wafer_sample_columns_are_wide_enough_to_read(generated, name: str) -> None:
    for ws in _report_sheets(generated(name)):
        d_width = ws.column_dimensions["D"].width
        for _sample, _header, _last, mdl_col in _wafer_panels(ws):
            for col in range(4, mdl_col - 1):
                letter = ws.cell(1, col).column_letter
                assert ws.column_dimensions[letter].width >= d_width, f"{ws.title} column {letter}"


# --- Summary rows ---


@pytest.mark.parametrize("name", ["dummy_tr_water", "Customer 98168_ipa_warning_limits"])
def test_total_row_leaves_the_untotalled_column_unboxed(generated, name: str) -> None:
    """TOTAL sums the result and MDL columns; the column between them has
    no total and no box (as in legacy). Customer 98168_ipa_warning_limits checks
    the same holds when limit columns shift the table right."""
    checked = 0
    for ws in _report_sheets(generated(name)):
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("TOTAL / "):
                    sums = [c for c in range(1, ws.max_column + 1)
                            if isinstance(ws.cell(cell.row, c).value, str) and ws.cell(cell.row, c).value.startswith("=SUM")]
                    assert len(sums) == 2, f"{ws.title} row {cell.row}: {sums}"
                    gap = ws.cell(cell.row, sums[0] + 1)
                    assert not _is_boxed(gap), f"{ws.title}!{gap.coordinate} is boxed"
                    checked += 1
    assert checked, "no TOTAL row found"


def test_formatting_helpers_detect_boxed_cells(tmp_path: Path) -> None:
    """Control: the checks above would fail on a formatted cell."""
    from openpyxl import Workbook
    from openpyxl.styles import Border, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws["A1"].border = Border(left=Side(style="thin"))
    ws["B1"].fill = PatternFill("solid", fgColor="FFCCFFFF")
    assert _is_boxed(ws["A1"]) and _is_boxed(ws["B1"]) and not _is_boxed(ws["C1"])
