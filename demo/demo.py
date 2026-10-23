"""Runnable demo with obviously fictional data: seeds a fresh database,
writes Testing Request workbooks by cell position (no real intake form),
and generates reports through the normal pipeline (cli.run -- no demo-only
code path).

    python demo/demo.py [--out demo_out]

Creates demo_out/demo.db, demo_out/requests/*.xlsx and
demo_out/reports/*.xlsx. Customers are "A".."Z"; chemicals are
IPA, HNO3, Acetone, H2O2 and PGMEA, plus "Water". A few customers show specific
features:

  A  no limits on file             B  specification limits
  C  action + warning limits       E  additional elements by another method
  F  Water limits                  G  a customer group; sites H and I use G's limits

The request writer mirrors the parser's cell map (slim-domain
tr_submission_form_parser / tr_sample_form_parser); it will be replaced
by the lab submission portal eventually.
"""

from __future__ import annotations

import argparse
import string
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path

import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from infrastructure.database import Base
from slim_domain.domain.analysis.analysis_repository import _AnalysisRow, _FormAnalysisRow
from slim_domain.domain.chemical.chemical_repository import _ChemicalRow, _FormChemicalRow, _KEDElementRow
from slim_domain.domain.customer.customer_repository import _CustomerRow, _FormCustomerRow
from slim_domain.domain.element.element_repository import _ElementRow
from slim_domain.domain.report_folder import report_folder_repository  # noqa: F401 -- registers its table
from slim_domain.domain.report_log import report_log_repository  # noqa: F401 -- registers its table
from slim_domain.domain.specification.specification_repository import _SpecificationRow

from slim_report_engine.reporting.chemical_water_report_builder import _UNCONDITIONALLY_SUPPORTED

_ELEMENTS = (
    "H Hydrogen,He Helium,Li Lithium,Be Beryllium,B Boron,C Carbon,N Nitrogen,O Oxygen,F Fluorine,Ne Neon,"
    "Na Sodium,Mg Magnesium,Al Aluminum,Si Silicon,P Phosphorus,S Sulfur,Cl Chlorine,Ar Argon,K Potassium,"
    "Ca Calcium,Sc Scandium,Ti Titanium,V Vanadium,Cr Chromium,Mn Manganese,Fe Iron,Co Cobalt,Ni Nickel,"
    "Cu Copper,Zn Zinc,Ga Gallium,Ge Germanium,As Arsenic,Se Selenium,Br Bromine,Kr Krypton,Rb Rubidium,"
    "Sr Strontium,Y Yttrium,Zr Zirconium,Nb Niobium,Mo Molybdenum,Tc Technetium,Ru Ruthenium,Rh Rhodium,"
    "Pd Palladium,Ag Silver,Cd Cadmium,In Indium,Sn Tin,Sb Antimony,Te Tellurium,I Iodine,Xe Xenon,"
    "Cs Cesium,Ba Barium,La Lanthanum,Ce Cerium,Pr Praseodymium,Nd Neodymium,Pm Promethium,Sm Samarium,"
    "Eu Europium,Gd Gadolinium,Tb Terbium,Dy Dysprosium,Ho Holmium,Er Erbium,Tm Thulium,Yb Ytterbium,"
    "Lu Lutetium,Hf Hafnium,Ta Tantalum,W Tungsten,Re Rhenium,Os Osmium,Ir Iridium,Pt Platinum,Au Gold,"
    "Hg Mercury,Tl Thallium,Pb Lead,Bi Bismuth,Po Polonium,At Astatine,Rn Radon,Fr Francium,Ra Radium,"
    "Ac Actinium,Th Thorium,Pa Protactinium,U Uranium"
)
# name -> (metals prep, ions prep)
_CHEMICALS = {
    "Water": ("Dilute and Shoot", ""),
    "IPA": ("Evaporation", ""),
    "HNO3": ("Dilute and Shoot", ""),
    "Acetone": ("Evaporation", "Evaporation"),
    "H2O2": ("Dilute and Shoot", ""),
    "PGMEA": ("Evaporation", ""),
}
_SITE_GROUP = {"H": "G", "I": "G"}  # site -> group


def seed(db_url: str) -> None:
    """Creates the schema and fills it with fictional reference data."""
    engine = create_engine(db_url)
    Base.metadata.create_all(engine)
    with Session(engine) as s:
        for i, pair in enumerate(_ELEMENTS.split(","), start=1):
            symbol, name = pair.split(" ", 1)
            s.add(_ElementRow(id=i, element_symbol=symbol, element_name=name))
        for i, name in enumerate(sorted(_UNCONDITIONALLY_SUPPORTED), start=1):
            s.add(_AnalysisRow(id=i, analysis_name=name, analysis_description=name))
            s.add(_FormAnalysisRow(form_name=name, analysis_id=i))
        chem_ids = {}
        for i, (name, (metals_prep, ions_prep)) in enumerate(_CHEMICALS.items(), start=1):
            s.add(_ChemicalRow(id=i, chemical_name=name, metals_prep=metals_prep, silicon_prep="",
                               ions_prep=ions_prep, metals_instrument="ICPMS", created_by="DEMO"))
            s.add(_FormChemicalRow(form_name=name, chemical_id=i, created_by="DEMO"))
            chem_ids[name] = i
        s.add(_KEDElementRow(chemical_id=chem_ids["HNO3"], element_id=26))  # Fe flagged KED
        cust_ids = {}
        for i, letter in enumerate(string.ascii_uppercase, start=1):
            s.add(_CustomerRow(
                id=i, customer_name=letter, street_address=f"{i} {letter} Street", city="Anytown",
                state="TX", postal_code="75000", country="USA",
                additional_elements_prep="Organic Dilute and Shoot" if letter == "E" else "",
                is_group=(letter == "G"), created_by="DEMO",
            ))
            cust_ids[letter] = i
        s.flush()
        for site, group in _SITE_GROUP.items():
            s.get(_CustomerRow, cust_ids[site]).parent_customer_id = cust_ids[group]
        for letter, cid in cust_ids.items():
            if letter != "G":
                s.add(_FormCustomerRow(form_customer=letter, form_address=f"{cid} {letter} Street",
                                       form_address_2="Anytown, TX 75000", customer_id=cid, created_by="DEMO"))
        next_id = 1

        def limits(customer: str, chemical: str, kind: str, label: str, values: dict[str, float]) -> None:
            nonlocal next_id
            for symbol, value in values.items():
                s.add(_SpecificationRow(
                    id=next_id, kind=kind, customer_id=cust_ids[customer], customer_name_raw=customer,
                    chemical_id=chem_ids[chemical], chemical_name_raw=chemical, element_symbol=symbol,
                    value=value, unit="ppb", category_label=label, parameter_text=symbol,
                ))
                next_id += 1

        limits("B", "IPA", "specification", "Specification", {"Al": 0.5, "Fe": 0.3, "Na": 1.0, "Zn": 0.2})
        limits("C", "HNO3", "specification", "Action Limit", {"Al": 0.4, "Fe": 0.2, "Ca": 0.75})
        limits("C", "HNO3", "warning_limit", "Warning Limit", {"Al": 0.3, "Fe": 0.1, "Ca": 0.5})
        limits("E", "IPA", "specification", "Specification", {"Sb": 0.1, "As": 0.1, "B": 1.0})
        limits("F", "Water", "specification", "Specification", {"Na": 5.0, "Si": 3.0, "Cl": 5.0})
        limits("G", "Acetone", "specification", "Specification", {"Al": 0.2, "Cu": 0.1, "Ni": 0.1})
        s.commit()


@dataclass
class DemoSample:
    sample_id: str
    analyses: list[str]
    chemical: str = "Water"  # ignored on a Water request
    additional_elements: str = ""
    processing_time: str = "Next Day"


@dataclass
class DemoRequest:
    customer: str  # "A".."Z"
    request_type: str  # "Chemical" or "Water"
    samples: list[DemoSample] = field(default_factory=list)
    received: date = field(default_factory=lambda: date.today() - timedelta(days=2))


def write_request(req: DemoRequest, path: Path) -> Path:
    """Writes a Testing Request workbook by cell position only -- the
    same cells the parser reads, nothing else."""
    chemical = req.request_type == "Chemical"
    cust_id = string.ascii_uppercase.index(req.customer) + 1
    proc_col = 11 if chemical else 14
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["D3"] = f"{req.request_type} Testing Request Form"
    ws["I4"] = req.received - timedelta(days=1)
    ws["P5" if chemical else "S5"] = req.received
    ws["P7" if chemical else "S7"] = "DM"  # technician initials
    ws["C10"], ws["C11"], ws["C12"] = req.customer, f"{cust_id} {req.customer} Street", "Anytown, TX 75000"
    ws["C13"], ws["B14"], ws["C14"] = "Demo Contact", "Phone :", "555-0100"
    ws["C15"] = "results@example.com"
    ws["P49" if chemical else "S51"] = "Rev Z"
    ws["H10"] = "PO-DEMO"
    ws["B20"] = "Sample ID"
    ws.cell(row=20, column=proc_col).value = "Processing Time"  # header ends here -> analyses stop before it
    row = 22
    for smp in req.samples:
        ws.cell(row=row, column=2).value = smp.sample_id
        if chemical:
            ws.cell(row=row, column=3).value = smp.chemical
        ws.cell(row=row, column=5).value = smp.additional_elements or None
        # Analyses go in columns 6.. (column 4 is unused here, 5 is additional elements).
        for col, name in enumerate(smp.analyses, start=6):
            if col >= proc_col:
                raise ValueError(f"too many analyses on {smp.sample_id} for a {req.request_type} request")
            ws.cell(row=row, column=col).value = name
        ws.cell(row=row, column=proc_col).value = smp.processing_time
        row += 1
    ws.cell(row=row - 1 + 11, column=2).value = "Total"  # footer sentinel: last sample = this row - 11
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


DEMO_REQUESTS = {
    "A plain": DemoRequest("A", "Chemical", [DemoSample("A-1", ["36 Elements", "5 Anions"], "IPA")]),
    "B limits": DemoRequest("B", "Chemical", [DemoSample("B-1", ["36 Elements"], "IPA", processing_time="Same Day RUSH")]),
    "C warning limits": DemoRequest("C", "Chemical", [DemoSample("C-1", ["36 Elements"], "HNO3")]),
    "E other method": DemoRequest("E", "Chemical", [DemoSample("E-1", ["36 Elements"], "IPA", "Sb, As, B")]),
    "F water": DemoRequest("F", "Water", [DemoSample("F-1", ["Anions", "Total Silicon"], additional_elements="Na"),
                                          DemoSample("F-2", ["TOC", "pH", "Conductivity"])]),
    "H group site": DemoRequest("H", "Chemical", [DemoSample("H-1", ["36 Elements"], "Acetone")]),
}


def run_demo(out: Path) -> list[Path]:
    from slim_report_engine.cli import build_services, run

    out.mkdir(parents=True, exist_ok=True)
    db = out / "demo.db"
    if db.exists():
        db.unlink()
    db_url = f"sqlite:///{db}"
    seed(db_url)
    services = build_services(db_url)
    stamp = date.today().strftime("%m%d%y")
    reports = []
    for label, req in DEMO_REQUESTS.items():
        request = write_request(req, out / "requests" / f"{stamp}-Demo {label} STD.xlsx")
        report = out / "reports" / f"{stamp}-Demo {label} Results.xlsx"
        report.parent.mkdir(parents=True, exist_ok=True)
        run(request, report, services=services)
        reports.append(report)
    return reports


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--out", type=Path, default=Path("demo_out"))
    args = parser.parse_args()
    for report in run_demo(args.out):
        print(f"wrote {report}")
